"""
document — 다양한 office/text 포맷 텍스트 추출 통합 모듈.

CLAUDE.md §2 (Raw-first / Loss-less): 추출 텍스트는 분석/검색용 derived data.
**원본 파일은 항상 별도 attachment 로 보존** (caller 책임 — backend/storage/local).

지원 포맷 (Phase 2.5):
- PDF  → backend.ingest.pdf._extract_pdf_text 재사용
- DOCX → python-docx
- PPTX → python-pptx
- TXT  → utf-8 / charset-normalizer 자동 감지
- MD   → utf-8

지원 안 함 (Phase 3+):
- DOC, PPT (구 binary 형식) — LibreOffice CLI 필요
- HWP/HWPX — pyhwp (한국 정부 문서)
- 이미지 — OCR (EasyOCR/Tesseract)

설계 — 진입점 2개:
- `guess_format(filename, mime_type)` → str — 포맷 식별
- `extract_text_from_bytes(data, filename, mime_type)` → DocumentExtract | None
  — 지원 안 하는 포맷이면 None (caller 는 attachment 만 저장)
"""

from __future__ import annotations

import io
import logging
import mimetypes
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

# PDF 재사용 — backend/ingest/pdf 의 검증된 함수.
from backend.ingest.pdf import _extract_pdf_text, _sanitize_text

logger = logging.getLogger(__name__)


@dataclass
class DocumentExtract:
    """포맷-agnostic 텍스트 추출 결과.

    `raw_text` 는 items.raw_content 로 들어감 — _sanitize_text 로 NUL byte 등 정제 완료.
    `title` 은 items.title 의 fallback (없으면 파일명).
    """
    raw_text: str
    title: str | None = None
    source_format: str = "unknown"        # "pdf" | "docx" | "pptx" | "txt" | "markdown"
    extractor: str = "unknown"            # "pypdf" | "pymupdf" | "python-docx" | "python-pptx" | "utf-8" | ...
    meta: dict[str, Any] = field(default_factory=dict)


# ──────────────────────────────────────────────────────────────
# Format 식별
# ──────────────────────────────────────────────────────────────


_EXT_TO_FORMAT: dict[str, str] = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".pptx": "pptx",
    ".txt": "txt",
    ".md": "markdown",
    ".markdown": "markdown",
    ".rst": "txt",          # rst 도 plain text 로 일단 (Phase 3+ 에 별 파서)
}


_MIME_TO_FORMAT: dict[str, str] = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
    "text/plain": "txt",
    "text/markdown": "markdown",
    "text/x-markdown": "markdown",
}


# Phase 3+ 에 추가 예정 — 지금은 attachment 만 저장
_KNOWN_UNSUPPORTED_FORMATS: set[str] = {
    "doc", "ppt", "hwp", "hwpx", "rtf",
    "image",  # png/jpg/gif/webp — caller 가 'image' 로 분류
}


def guess_format(
    filename: str | None = None, mime_type: str | None = None,
) -> str:
    """mime_type → 확장자 → "unknown" 순으로 포맷 식별.

    이미지 (mime_type 이 image/* 로 시작) 는 "image" 로 분류 — extract 는 None.
    """
    if mime_type:
        mt = mime_type.lower().split(";")[0].strip()
        if mt in _MIME_TO_FORMAT:
            return _MIME_TO_FORMAT[mt]
        if mt.startswith("image/"):
            return "image"

    if filename:
        ext = Path(filename).suffix.lower()
        if ext in _EXT_TO_FORMAT:
            return _EXT_TO_FORMAT[ext]
        # 구 binary office 명시적 표시 (Phase 3+ 에서 LibreOffice 로 처리)
        if ext == ".doc":
            return "doc"
        if ext == ".ppt":
            return "ppt"
        if ext in {".hwp", ".hwpx"}:
            return "hwp"
        if ext in {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tiff"}:
            return "image"

    return "unknown"


def is_supported(fmt: str) -> bool:
    """guess_format 결과가 텍스트 추출 가능한지."""
    return fmt in {"pdf", "docx", "pptx", "txt", "markdown"}


# ──────────────────────────────────────────────────────────────
# 포맷별 추출
# ──────────────────────────────────────────────────────────────


def _extract_docx(data: bytes) -> tuple[str, str | None, dict[str, Any]]:
    """python-docx — paragraph 텍스트 join."""
    from docx import Document   # type: ignore[import-not-found]

    doc = Document(io.BytesIO(data))
    paragraphs = [p.text for p in doc.paragraphs if p.text]
    # 표 안의 텍스트도 가져옴 (실무 docx 는 표가 많음)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                txt = cell.text
                if txt:
                    paragraphs.append(txt)
    body = "\n\n".join(paragraphs)

    # core_properties — title 우선, 없으면 None (caller 가 filename 으로 fallback)
    title = None
    try:
        cp = doc.core_properties
        if cp.title:
            title = cp.title.strip() or None
    except Exception:  # noqa: BLE001
        pass

    meta: dict[str, Any] = {
        "paragraph_count": len(paragraphs),
        "table_count": len(doc.tables),
    }
    return body, title, meta


def _extract_pptx(data: bytes) -> tuple[str, str | None, dict[str, Any]]:
    """python-pptx — 슬라이드별 텍스트 + notes 까지."""
    from pptx import Presentation   # type: ignore[import-not-found]

    pres = Presentation(io.BytesIO(data))
    slides_text: list[str] = []
    for i, slide in enumerate(pres.slides, start=1):
        parts: list[str] = [f"=== Slide {i} ==="]
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    txt = "".join(r.text for r in p.runs).strip()
                    if txt:
                        parts.append(txt)
        # 발표자 노트 — 학습 데이터로 가치 큼
        if slide.has_notes_slide:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()
            if notes_text:
                parts.append(f"[Notes] {notes_text}")
        slides_text.append("\n".join(parts))
    body = "\n\n".join(slides_text)

    # title slide 의 첫 텍스트를 title 로
    title = None
    if pres.slides:
        first = pres.slides[0]
        for shape in first.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                title = shape.text_frame.text.strip().split("\n")[0][:200]
                break

    meta: dict[str, Any] = {"slide_count": len(pres.slides)}
    return body, title, meta


def _extract_text_file(data: bytes, fmt: str) -> tuple[str, dict[str, Any]]:
    """TXT/MD — 인코딩 자동 감지.

    순서:
    1. utf-8 (가장 흔함, 거의 모든 모던 환경)
    2. cp949 / euc-kr (한국 윈도우 메모장 / 옛 한글 텍스트 — charset-normalizer
       는 짧은 텍스트로 big5/cp949 구분 못 함, 한국 사용자 환경 우선)
    3. charset-normalizer (longer text 의 일반 자동 감지)
    4. 최후: utf-8 errors='replace' — loss-less 위반이지만 raw 는 attachment 로 보존됨
    """
    # 1. utf-8 — 가장 흔함
    try:
        return data.decode("utf-8"), {"encoding": "utf-8"}
    except UnicodeDecodeError:
        pass

    # 2. 한국어 인코딩 우선 시도 (사용자 한국 환경, 옛 메모장 파일)
    for enc in ("cp949", "euc-kr"):
        try:
            return data.decode(enc), {"encoding": enc}
        except UnicodeDecodeError:
            continue

    # 3. charset-normalizer 일반 자동 감지
    try:
        from charset_normalizer import from_bytes   # type: ignore[import-not-found]

        result = from_bytes(data).best()
        if result is not None:
            return str(result), {"encoding": result.encoding or "auto"}
    except Exception:  # noqa: BLE001
        pass

    # 4. 최후 fallback
    return data.decode("utf-8", errors="replace"), {"encoding": "utf-8-replace"}


# ──────────────────────────────────────────────────────────────
# 통합 진입점
# ──────────────────────────────────────────────────────────────


def extract_text_from_bytes(
    data: bytes,
    *,
    filename: str | None = None,
    mime_type: str | None = None,
) -> DocumentExtract | None:
    """파일 bytes → DocumentExtract.

    Args:
        data: 파일 raw bytes
        filename: 원 파일명 (확장자 + title fallback 용)
        mime_type: 알면 우선 사용 (없으면 filename 으로 추정)

    Returns:
        지원 포맷이면 DocumentExtract, 아니면 None (caller 가 attachment 만 저장).

    실패 시:
        파싱 예외 발생하면 logger.warning 후 None 반환 — caller 가 attachment 만
        저장하도록 graceful degrade.
    """
    if not data:
        return None

    # mime_type 모르면 filename 으로 추정
    if not mime_type and filename:
        guessed_mt, _ = mimetypes.guess_type(filename)
        mime_type = guessed_mt

    fmt = guess_format(filename, mime_type)

    if fmt == "pdf":
        try:
            text_out, info = _extract_pdf_text(data)
        except Exception as e:  # noqa: BLE001
            logger.warning(
                "PDF 추출 실패 (file=%s, %s: %s)",
                filename, type(e).__name__, e,
            )
            return None
        if not text_out.strip():
            return None
        return DocumentExtract(
            raw_text=text_out,
            title=_filename_to_title(filename),
            source_format="pdf",
            extractor=info.get("extractor", "pdf"),
            meta=info,
        )

    if fmt == "docx":
        try:
            body, title, meta = _extract_docx(data)
        except Exception as e:  # noqa: BLE001
            logger.warning(
                "DOCX 추출 실패 (file=%s, %s: %s)",
                filename, type(e).__name__, e,
            )
            return None
        if not body.strip():
            return None
        return DocumentExtract(
            raw_text=_sanitize_text(body),
            title=title or _filename_to_title(filename),
            source_format="docx",
            extractor="python-docx",
            meta=meta,
        )

    if fmt == "pptx":
        try:
            body, title, meta = _extract_pptx(data)
        except Exception as e:  # noqa: BLE001
            logger.warning(
                "PPTX 추출 실패 (file=%s, %s: %s)",
                filename, type(e).__name__, e,
            )
            return None
        if not body.strip():
            return None
        return DocumentExtract(
            raw_text=_sanitize_text(body),
            title=title or _filename_to_title(filename),
            source_format="pptx",
            extractor="python-pptx",
            meta=meta,
        )

    if fmt in {"txt", "markdown"}:
        try:
            body, info = _extract_text_file(data, fmt)
        except Exception as e:  # noqa: BLE001
            logger.warning(
                "TXT/MD 추출 실패 (file=%s, %s: %s)",
                filename, type(e).__name__, e,
            )
            return None
        if not body.strip():
            return None
        # MD 의 경우 첫 # heading 을 title 로 시도
        title = _filename_to_title(filename)
        if fmt == "markdown":
            for line in body.splitlines():
                line = line.strip()
                if line.startswith("# "):
                    title = line[2:].strip() or title
                    break
        return DocumentExtract(
            raw_text=_sanitize_text(body),
            title=title,
            source_format=fmt,
            extractor=info.get("encoding", "utf-8"),
            meta=info,
        )

    # unsupported (doc/ppt/hwp/image/unknown) — attachment 만 저장하도록 None
    return None


def _filename_to_title(filename: str | None) -> str | None:
    """확장자 제거 + 너무 길면 cut. 파일명만으로 부족하면 None."""
    if not filename:
        return None
    name = Path(filename).stem.strip()
    if not name:
        return None
    return name[:200]
